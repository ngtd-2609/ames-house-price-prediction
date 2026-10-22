#!/usr/bin/env python
"""Independent consistency checker driven by current machine-readable evidence.

The checker never hard-codes a test count or coverage percentage.  It derives
truth from the canonical model/summary plus the latest structured audit files.
"""
from __future__ import annotations

from hashlib import sha256
import json
from pathlib import Path
import re
import sys
import zipfile
from datetime import datetime
from xml.etree import ElementTree

try:
    import nbformat
except ImportError:
    class _NotebookCell(dict):
        __getattr__ = dict.__getitem__

    class _NotebookFormatFallback:
        @staticmethod
        def read(path, as_version=4):
            del as_version
            payload = json.loads(Path(path).read_text(encoding="utf-8"))
            cells = []
            for cell in payload.get("cells", []):
                normalized = dict(cell)
                if isinstance(normalized.get("source"), list):
                    normalized["source"] = "".join(normalized["source"])
                cells.append(_NotebookCell(normalized))
            return _NotebookCell(
                {**payload, "cells": cells}
            )

    nbformat = _NotebookFormatFallback()
from docx import Document
try:
    from pptx import Presentation
except ImportError:
    Presentation = None

ROOT = Path(__file__).resolve().parents[1]
AUDIT = ROOT / "outputs" / "audit"
ERRORS: list[str] = []
STATUS: dict[str, bool] = {}


def load_json(rel: str) -> dict:
    path = ROOT / rel
    if not path.exists():
        ERRORS.append(f"MISSING {rel}")
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        ERRORS.append(f"INVALID_JSON {rel}: {exc}")
        return {}


def file_sha256(path: Path) -> str:
    h = sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def docx_text(path: Path) -> str:
    if not path.exists():
        return ""
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def pptx_text(path: Path) -> str:
    if not path.exists():
        return ""
    if Presentation is None:
        chunks: list[str] = []
        with zipfile.ZipFile(path) as archive:
            names = [
                name for name in archive.namelist()
                if (name.startswith("ppt/slides/slide") or name.startswith("ppt/notesSlides/notesSlide"))
                and name.endswith(".xml")
            ]
            for name in names:
                root = ElementTree.fromstring(archive.read(name))
                chunks.extend(node.text or "" for node in root.iter() if node.tag.endswith("}t"))
        return "\n".join(chunks)
    prs = Presentation(path)
    chunks: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                chunks.append(shape.text)
            if shape.has_table:
                for row in shape.table.rows:
                    chunks.extend(cell.text for cell in row.cells)
        try:
            chunks.append(slide.notes_slide.notes_text_frame.text)
        except Exception:
            pass
    return "\n".join(chunks)


def mark(name: str, ok: bool, detail: str = "") -> None:
    STATUS[name] = bool(ok)
    if not ok:
        ERRORS.append(f"{name}: {detail or 'FAIL'}")


def metric_tokens(metric_dict: dict) -> list[str]:
    out=[]
    for key,value in metric_dict.items():
        if isinstance(value,(int,float)):
            out.extend([f"{float(value):,.2f}", f"{float(value):.4f}"])
    return out


summary = load_json("outputs/summary.json")
metadata = load_json("models/model_metadata.json")
policy = load_json("outputs/model_selection_policy.json")
test_summary = load_json("outputs/audit/test_summary.json")
coverage = load_json("outputs/audit/coverage.json")
notebook_summary = load_json("outputs/notebook_summary.json")

model_path = ROOT / summary.get("artifact", {}).get("path", "models/house_price_pipeline.joblib")
actual_model_hash = file_sha256(model_path) if model_path.exists() else ""
expected_model = summary.get("selected_model")
expected_params = summary.get("selected_params")
expected_test = summary.get("test_metrics", {})
expected_val = summary.get("validation_metrics", {})
expected_hash = summary.get("artifact", {}).get("model_sha256")

mark("MODEL_NAME", bool(expected_model) and metadata.get("selected_model") == expected_model and policy.get("locked_model") == expected_model,
     f"summary={expected_model}, metadata={metadata.get('selected_model')}, policy={policy.get('locked_model')}")
mark("HYPERPARAMETERS", bool(expected_params) and metadata.get("selected_params") == expected_params and policy.get("locked_params") == expected_params,
     "selected_params disagree")
mark("VALIDATION_METRICS", metadata.get("validation_metrics") == expected_val, "metadata validation metrics disagree")
mark("TEST_METRICS", metadata.get("test_metrics") == expected_test, "metadata test metrics disagree")
mark("ARTIFACT_HASH", bool(actual_model_hash) and expected_hash == actual_model_hash and metadata.get("model_sha256") == actual_model_hash,
     f"actual={actual_model_hash}, summary={expected_hash}, metadata={metadata.get('model_sha256')}")

# Policy JSON must equal executable defaults for all executable configuration fields.
sys.path.insert(0, str(ROOT / "src"))
try:
    from house_price.selection import ModelSelectionPolicy
    defaults = ModelSelectionPolicy()
    policy_matches = (
        policy.get("primary_metric") == defaults.primary_metric
        and policy.get("validation_tolerance") == defaults.validation_tolerance
        and policy.get("ridge_tie_break_requires_cv_not_worse") == defaults.ridge_tie_break_requires_cv_not_worse
        and policy.get("uses_test") is defaults.uses_test
        and policy.get("candidate_filter") == defaults.candidate_filter
        and policy.get("tie_breaker_order") == list(defaults.tie_breaker_order)
    )
except Exception as exc:
    policy_matches=False; ERRORS.append(f"POLICY_IMPORT: {exc}")
mark("SELECTION_POLICY_CODE_VS_JSON", policy_matches, "executable defaults differ from JSON")

# Structured test/coverage evidence, with freshness against source/tests.
passed = test_summary.get("tests_passed")
failed = test_summary.get("tests_failed")
coverage_value = coverage.get("total_coverage")
scope = coverage.get("scope")
mark("TEST_COUNT", isinstance(passed, int) and passed > 0 and failed == 0, f"passed={passed}, failed={failed}")
mark("SOURCE_COVERAGE", isinstance(coverage_value, (int,float)) and coverage_value >= 95 and scope == ["src/house_price", "backend"],
     f"coverage={coverage_value}, scope={scope}")
important = ["src/house_price/tuning.py", "src/house_price/reporting.py", "src/house_price/validation.py", "src/house_price/artifacts.py", "backend/service.py"]
module_ok = all(coverage.get("files", {}).get(rel, {}).get("coverage_pct", 0) > 0 for rel in important)
mark("IMPORTANT_MODULE_COVERAGE", module_ok, "an important module has zero/missing coverage")

try:
    evidence_time = datetime.fromisoformat(str(coverage["generated_at_utc"]).replace("Z", "+00:00")).timestamp()
    newest_code = max(p.stat().st_mtime for folder in [ROOT/"src", ROOT/"backend", ROOT/"tests"] for p in folder.rglob("*.py"))
    fresh_evidence = evidence_time + 2 >= newest_code
except Exception:
    fresh_evidence = False
mark("AUDIT_EVIDENCE_FRESHNESS", fresh_evidence, "coverage/test evidence predates source/tests")

# Notebook source truth: actual counts/status + model lock ordering.
nb_ok=True; nb_count_ok=True; protocol_ok=True
for prefix, name in [("detailed", "House_Price_Prediction_Detailed.ipynb"), ("basic", "House_Price_Prediction_Basic.ipynb")]:
    path=ROOT/"notebooks"/name
    if not path.exists():
        nb_ok=nb_count_ok=protocol_ok=False; continue
    nb=nbformat.read(path, as_version=4)
    errors=[o for c in nb.cells if c.cell_type=="code" for o in c.get("outputs",[]) if o.get("output_type")=="error"]
    unexecuted=[c for c in nb.cells if c.cell_type=="code" and c.get("execution_count") is None]
    nb_ok &= not errors and not unexecuted and notebook_summary.get(f"{prefix}_status") == "PASS"
    nb_count_ok &= notebook_summary.get(f"{prefix}_cells") == len(nb.cells)
    sources=[c.source for c in nb.cells]
    lock_idx=next((i for i,s in enumerate(sources) if "MODEL LOCKED" in s), -1)
    test_load_idx=next((i for i,s in enumerate(sources) if re.search(r'\b(test|te)\s*=\s*load_house_data\("house_test_10\.csv"\)', s)), -1)
    ytest_idx=next((i for i,s in enumerate(sources) if re.search(r'\by_test\b|\byte\b', s) and "split_xy" in s), -1)
    protocol_ok &= lock_idx >= 0 and test_load_idx > lock_idx and ytest_idx >= test_load_idx
mark("NOTEBOOK_STATUS", nb_ok, "notebook execution error/unexecuted code")
mark("NOTEBOOK_CELL_COUNTS", nb_count_ok, "notebook_summary cell count is stale")
mark("NOTEBOOK_MODEL_LOCK_PROTOCOL", protocol_ok, "test/y_test appears before MODEL LOCKED")

# Human artifacts. Docs that mention evidence must match current evidence; stale historic final claims are forbidden.
coverage_display = f"{float(coverage.get('display_coverage', 0)):.1f}%"
passed_token = str(passed) if passed is not None else ""
rmse_token = f"{float(expected_test.get('RMSE', 0)):,.2f}"
val_rmse_token = f"{float(expected_val.get('RMSE', 0)):,.2f}"
hash16 = actual_model_hash[:16]
feature_count = str(summary.get("dataset", {}).get("input_features_raw", ""))
split_tokens = [str(summary.get("dataset", {}).get(k, [""])[0]) for k in ["train_shape", "validation_shape", "test_shape"]]

md_paths = [
    "README.md", "DATA_CARD.md", "MODEL_CARD.md", "REPRODUCIBILITY.md",
    "docs/ARCHITECTURE.md", "docs/DEMO_GUIDE.md", "docs/DEPLOYMENT.md", "DEFENSE_QA.md",
]
texts={rel:(ROOT/rel).read_text(encoding="utf-8",errors="ignore") if (ROOT/rel).exists() else "" for rel in md_paths}
stale_patterns=[r"\b56 passed\b", r"\bcoverage\s*[:=]?\s*96%", r"\b47 tests?\b", r"\b91%\b"]
stale=[]
for rel,text in texts.items():
    if not text:
        stale.append(f"MISSING {rel}"); continue
    # Historical/baseline evidence is legitimate when explicitly labelled as such.
    # Flag an old number only when it appears as an unqualified current/final claim.
    for line in text.splitlines():
        low=line.lower()
        historical=any(tag in low for tag in ("baseline", "before", "trước khi sửa", "historical", "previous audit"))
        if historical:
            continue
        for pat in stale_patterns:
            if re.search(pat,line,flags=re.I): stale.append(f"{rel}: {pat}")
    if re.search(r"(?:source\s+)?coverage",text,flags=re.I) and coverage_display not in text:
        stale.append(f"{rel}: coverage claim does not contain {coverage_display}")
mark("STALE_REFERENCES", not stale, "; ".join(stale[:12]))

for rel in ["README.md", "MODEL_CARD.md"]:
    text=texts.get(rel,"")
    required=[expected_model, rmse_token]
    if rel != "MODEL_CARD.md": required += [passed_token, coverage_display]
    ok=all(str(token) in text for token in required if token)
    mark(rel.upper().replace(".","_"), ok, f"missing current canonical tokens in {rel}")

report_path=ROOT/"reports"/"Bao_cao_Hoc_may.docx"
report=docx_text(report_path)
report_required=[expected_model, rmse_token, val_rmse_token, passed_token, coverage_display, hash16, feature_count]
split_ok=all(t in report or f"{int(t):,}".replace(",", ".") in report for t in split_tokens if t)
mark("DOCX", bool(report) and split_ok and all(t in report for t in report_required if t), "report missing canonical model/metrics/tests/coverage/hash/features/splits")

ppt_path=ROOT/"slides"/"Slide_Hoc_may.pptx"
ppt=pptx_text(ppt_path)
ppt=ppt.replace(coverage_display.replace(".", ","), coverage_display)
ppt_required=[expected_model, rmse_token, passed_token, coverage_display, "Raw input", "Feature engineering", "ColumnTransformer"]
mark("PPTX", bool(ppt) and all(t in ppt for t in ppt_required if t), "slides missing canonical testing/model/architecture evidence")

for key in [
    "MODEL_NAME","HYPERPARAMETERS","VALIDATION_METRICS","TEST_METRICS","ARTIFACT_HASH","TEST_COUNT","SOURCE_COVERAGE",
    "IMPORTANT_MODULE_COVERAGE","AUDIT_EVIDENCE_FRESHNESS","NOTEBOOK_STATUS","NOTEBOOK_CELL_COUNTS","NOTEBOOK_MODEL_LOCK_PROTOCOL",
    "SELECTION_POLICY_CODE_VS_JSON","README_MD","MODEL_CARD_MD","DOCX","PPTX","STALE_REFERENCES",
]:
    print(f"{key}: {'PASS' if STATUS.get(key) else 'FAIL'}")

if ERRORS:
    print("\nDETAILS:")
    for e in ERRORS:
        print("-",e)
    raise SystemExit(1)
